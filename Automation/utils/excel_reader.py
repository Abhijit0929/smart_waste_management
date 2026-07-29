from openpyxl import load_workbook


class ExcelReader:

    def __init__(self, file_path):
        self.workbook = load_workbook(file_path)

    def get_data(self, sheet_name):

        sheet = self.workbook[sheet_name]

        rows = list(sheet.iter_rows(values_only=True))

        headers = rows[0]

        data = []

        for row in rows[1:]:

            record = {}

            for header, value in zip(headers, row):
                record[header] = value

            data.append(record)

        return data